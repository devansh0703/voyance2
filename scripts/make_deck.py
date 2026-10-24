"""Voyance II investor/operator deck — built with python-pptx per the pptx skill.

Design system (matches the product's light terminal theme):
  BG_LIGHT #F4F5F1 (content) / BG_DARK #0B0D0C (cover+closing) · PRIMARY #0E6E5C ·
  ACCENT #FFB454 (amber) · TEXT #17201B · MUTED #5D675F
Fonts: Arial (everything) + Consolas (data figures) — both ship with Office.
Motif: hairline separators + oversized mono figures. No stripes, no underlines.
"""
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

BG_DARK = RGBColor(0x0B, 0x0D, 0x0C)
BG = RGBColor(0xF4, 0xF5, 0xF1)
PRIMARY = RGBColor(0x0E, 0x6E, 0x5C)
ACCENT = RGBColor(0xFF, 0xB4, 0x54)
TEXT = RGBColor(0x17, 0x20, 0x1B)
MUTED = RGBColor(0x5D, 0x67, 0x5F)
PANEL = RGBColor(0xFC, 0xFC, 0xFA)
RULE = RGBColor(0xD7, 0xDA, 0xD2)
BAD = RGBColor(0xC2, 0x40, 0x2E)

SANS = "Arial"
MONO = "Consolas"

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]


def slide(dark=False):
    s = prs.slides.add_slide(BLANK)
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = BG_DARK if dark else BG
    return s


def box(s, x, y, w, h, text, size, color, bold=False, font=SANS,
        align=PP_ALIGN.LEFT, spacing=1.0):
    tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    p = tf.paragraphs[0]
    p.alignment = align
    p.line_spacing = spacing
    r = p.add_run()
    r.text = text
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.name = font
    r.font.color.rgb = color
    return tb


def hairline(s, x, y, w, color=RULE):
    from pptx.enum.shapes import MSO_SHAPE
    ln = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Pt(1))
    ln.fill.solid()
    ln.fill.fore_color.rgb = color
    ln.line.fill.background()
    return ln


def kicker(s, text, dark=False, x=0.6, y=0.55):
    box(s, x, y, 10, 0.3, text.upper(), 11, MUTED if not dark else RGBColor(0x8B, 0x94, 0x8C),
        bold=True, font=MONO)


def title(s, text, dark=False, y=0.95, size=32, x=0.6, w=12.1):
    box(s, x, y, w, 0.9, text, size, TEXT if not dark else RGBColor(0xF4, 0xF5, 0xF1), bold=True)


def source(s, text, dark=False):
    box(s, 0.6, 7.05, 12, 0.3, text, 10, MUTED if not dark else RGBColor(0x8B, 0x94, 0x8C), font=MONO)


def bignum(s, x, y, value, label, color=PRIMARY, vsize=54):
    box(s, x, y, 3.6, 1.0, value, vsize, color, bold=True, font=MONO)
    box(s, x, y + 1.02, 3.6, 0.5, label.upper(), 11, MUTED, bold=True, font=MONO)


# ── 1 · COVER (dark) ─────────────────────────────────────────────
s = slide(dark=True)
box(s, 0.7, 0.7, 8, 0.4, "OPEN SOURCE · LIVE MARITIME INTELLIGENCE", 13, ACCENT, bold=True, font=MONO)
box(s, 0.7, 1.7, 11.5, 1.6, "Voyance II", 76, RGBColor(0xF4, 0xF5, 0xF1), bold=True)
box(s, 0.7, 3.15, 10.5, 0.9,
    "Know where your ETA model fails — before the storm does.",
    24, RGBColor(0xE7, 0xEC, 0xE6))
hairline(s, 0.7, 4.35, 11.9, RGBColor(0x2A, 0x2F, 0x2B))
for i, (v, l) in enumerate([("4.49 h", "TEST MAE · GPU MODEL"),
                             ("83%", "P10–P90 COVERAGE"),
                             ("8,020", "REAL VOYAGES MODELED"),
                             ("25+", "OPEN DATA SOURCES")]):
    box(s, 0.7 + i * 3.05, 4.7, 2.9, 0.7, v, 30, ACCENT, bold=True, font=MONO)
    box(s, 0.7 + i * 3.05, 5.35, 2.9, 0.4, l, 10, RGBColor(0x8B, 0x94, 0x8C), bold=True, font=MONO)
source(s, "github.com/devansh0703/voyance2 · Apache-2.0 · every number reproducible from pipeline artifacts", dark=True)

# ── 2 · PROBLEM ─────────────────────────────────────────────────
s = slide()
kicker(s, "the problem · idealab idea #586938")
title(s, "Operators judge AI models from vendor claims.")
box(s, 0.6, 2.0, 6.6, 3.6,
    "A fleet team trusts a vendor's arrival estimate but cannot see where it "
    "fails in bad weather or certain waters.\n\nWindward, PortXchange, WHIZTEC "
    "sell closed operational intelligence. Nobody ships an open, served, "
    "weather-conditioned model with honest error decomposition.",
    19, TEXT, spacing=1.25)
bignum(s, 8.1, 2.2, "$100M", "wedge market (model estimate)")
bignum(s, 8.1, 4.1, "$4.3B", "wider maritime AI market")
source(s, "Source: idealab.ai/idea-details/586938 market analysis (Grand View Research 2024 proxy)")

# ── 3 · WHAT IT IS ──────────────────────────────────────────────
s = slide()
kicker(s, "the product")
title(s, "A live maritime intelligence terminal, open end to end")
steps = [("INGEST", "NOAA MarineCadastre AIS\n+ UN/LOCODE ports"),
         ("VOYAGEIZE", "port-buffer segmentation\nplausibility filters"),
         ("WEATHERIZE", "ERA5 wind · waves · swell\nhourly grid join"),
         ("SERVE", "GPU quantile ETA\nFastAPI /v1 /v2"),
         ("WATCH", "drift PSI/KS\nEvidently + alerts")]
x = 0.6
for i, (h, d) in enumerate(steps):
    box(s, x, 2.15, 2.2, 0.35, f"0{i+1}", 22, ACCENT, bold=True, font=MONO)
    box(s, x, 2.62, 2.25, 0.35, h, 13, PRIMARY, bold=True, font=MONO)
    box(s, x, 3.05, 2.25, 1.1, d, 12.5, TEXT, spacing=1.15)
    if i < 4:
        box(s, x + 2.18, 2.2, 0.3, 0.4, "→", 18, RULE, bold=True)
    x += 2.45
hairline(s, 0.6, 4.55, 12.1)
box(s, 0.6, 4.8, 12.0, 1.0,
    "Every layer is open data and open code: NOAA/CC0 · Open-Meteo/CC BY 4.0 · OSM/ODbL · "
    "Digitraffic/CC BY 4.0 · EMODnet · USGS · OBIS/GBIF. Zero signups, zero fabricated numbers.",
    15, MUTED, spacing=1.3)

# ── 4 · MODEL BENCHMARK (native chart) ──────────────────────────
s = slide()
kicker(s, "the model · trained on the owner's RTX 4060 (CUDA)")
title(s, "Beats the speed prior by 16% — and stays calibrated")
cd = CategoryChartData()
cd.categories = ["XGBoost GPU (served)", "Speed-prior baseline", "LightGBM CPU"]
cd.add_series("Test MAE (h)", (4.49, 5.35, 5.80))
gf = s.shapes.add_chart(XL_CHART_TYPE.BAR_CLUSTERED, Inches(0.6), Inches(2.0),
                        Inches(7.2), Inches(4.4), cd)
ch = gf.chart
ch.has_legend = False
pl = ch.plots[0]
pl.gap_width = 60
for i, pt in enumerate(pl.series[0].points):
    pt.format.fill.solid()
    pt.format.fill.fore_color.rgb = [PRIMARY, RULE, RGBColor(0x9C, 0xC5, 0xB8)][i]
ch.value_axis.has_major_gridlines = False
box(s, 8.2, 2.2, 4.5, 0.8, "4.49 h", 54, PRIMARY, bold=True, font=MONO)
box(s, 8.2, 3.25, 4.5, 0.4, "TEST MAE — JULY 2024 HOLDOUT", 11, MUTED, bold=True, font=MONO)
hairline(s, 8.2, 3.9, 4.3)
box(s, 8.2, 4.1, 4.6, 2.2,
    "· Quantile P10/P50/P90, conformal-calibrated (CQR, Romano 2019)\n"
    "· 83% empirical coverage vs 80% nominal\n"
    "· Physics-informed ratio features — trees can't divide\n"
    "· 18,040 train / 12,924 test rows, voyage-disjoint split",
    13, TEXT, spacing=1.35)
source(s, "Source: reports/eval.json (generated by pipeline/train.py — not hand-entered)")

# ── 5 · PREDICTION SURFACE ──────────────────────────────────────
s = slide()
kicker(s, "the product is the prediction, laid out")
title(s, "Not a point estimate — a full voyage surface")
box(s, 0.6, 2.05, 5.9, 3.4,
    "· Actual track + remaining leg on real chart geography\n"
    "· Drift cone sized by the P10–P90 lateral spread\n"
    "· Historical analog voyages drawn dashed — same port, same ship type\n"
    "· ETA quantile cards with distance-to-go and fix count\n"
    "· One GeoJSON endpoint: /v2/predict/route/{voyage}",
    16, TEXT, spacing=1.5)
bignum(s, 7.4, 2.3, "P10/P50/P90", "conformal ETA band", vsize=30)
hairline(s, 7.4, 3.75, 5.2)
bignum(s, 7.4, 4.0, "±2.73 h", "CQR interval adjustment", vsize=44, color=ACCENT)
source(s, "Live at http://127.0.0.1:8010/predict · GeoJSON served by XGBoost quantile models")

# ── 6 · LIVE OPS ─────────────────────────────────────────────────
s = slide()
kicker(s, "live operations · no signup feeds")
title(s, "Real-time fleet picture with collision watch")
rows = [("LIVE AIS", "300 vessels · Digitraffic Finland · 20 s refresh"),
        ("COLLISION", "DCPA/TCPA solver · Goodwin-Stanton ellipses · fuzzy CRI"),
        ("TRACKING", "EKF + UKF filters · Douglas-Peucker · Fréchet residuals"),
        ("OVERLAYS", "MPA polygons · quakes · radar loop · SAR chips · harbors"),
        ("QUALITY", "IQR/z-score audit · upstream health · event ticker")]
y = 2.1
for h, d in rows:
    box(s, 0.6, y, 2.3, 0.4, h, 14, PRIMARY, bold=True, font=MONO)
    box(s, 3.0, y, 9.6, 0.4, d, 15, TEXT)
    hairline(s, 0.6, y + 0.5, 12.1)
    y += 0.82
source(s, "Encounter view draws safety-domain ellipses + CPA extension lines on the chart — /live")

# ── 7 · AI ANALYST ──────────────────────────────────────────────
s = slide()
kicker(s, "explainable by construction")
title(s, "An AI analyst that cites real voyages")
box(s, 0.6, 2.0, 6.4, 4.2,
    "\u201cWhy will this route drift tomorrow for this ship type?\u201d\n\n"
    "The answer is generated only from a context packet: live forecast + waves at "
    "the vessel, the model's tested MAE and top features, and the actual historical "
    "analog voyages — each named, each replayable on the map.\n\n"
    "Provider chain: OPENCODE_API_KEY first, chat-verified Groq/OpenRouter fallback "
    "— every response states who served it.",
    17, TEXT, spacing=1.35)
box(s, 7.5, 2.15, 5.2, 0.4, "PARSED OUTPUT", 12, MUTED, bold=True, font=MONO)
for i, sec in enumerate(["## Headline", "## Why drift is expected", "## Reference backing", "## What to watch"]):
    box(s, 7.5, 2.6 + i * 0.62, 5.2, 0.45, sec, 16, PRIMARY if i else ACCENT, bold=True, font=MONO)
hairline(s, 7.5, 5.15, 5.2)
box(s, 7.5, 5.35, 5.2, 0.9, "references = real voyage_ids\nreplayable in one click", 13, MUTED, font=MONO, spacing=1.3)
source(s, "POST /v2/ai/explain · answers forbidden from inventing numbers outside the context packet")

# ── 8 · DRIFT MONITORING ────────────────────────────────────────
s = slide()
kicker(s, "production honesty")
title(s, "Drift monitoring with alerts, not vibes")
cd = CategoryChartData()
cd.categories = ["Apr 2024", "Jul 2024"]
cd.add_series("Slice MAE ratio vs reference", (1.33, 1.0))
gf = s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(0.6), Inches(2.0),
                        Inches(5.6), Inches(4.2), cd)
ch = gf.chart
ch.has_legend = False
for i, pt in enumerate(ch.plots[0].series[0].points):
    pt.format.fill.solid()
    pt.format.fill.fore_color.rgb = [ACCENT, PRIMARY][i]
ch.value_axis.has_major_gridlines = False
box(s, 6.8, 2.1, 6.0, 3.9,
    "· Evidently DataDrift per feature: PSI + KS\n"
    "· Alert rule: PSI > 0.2 or MAE ratio > 1.25\n"
    "· April slice drifted on distance/SOG mix — caught,\n   explained, surfaced as a chip on the dashboard\n"
    "· Machine-readable alerts at /v1/drift",
    16, TEXT, spacing=1.5)
source(s, "Source: reports/drift_summary.json + Evidently HTML reports in reports/drift/")

# ── 9 · WHY IT WINS ─────────────────────────────────────────────
s = slide()
kicker(s, "positioning")
title(s, "The open slot in the landscape")
cols = [("CLOSED VENDORS", "Windward · PortXchange\nWHIZTEC · Broadpath\nKairon", "Operational visibility,\nnot independent error audit"),
        ("OPEN BENCHMARKS", "WUT-TrajData", "Dataset + baselines,\nnot a served product"),
        ("VOYANCE II", "Apache-2.0 · self-host", "Served quantile model +\nscorecard + AI analyst")]
x = 0.6
for i, (h, who, what) in enumerate(cols):
    hl = i == 2
    box(s, x, 2.1, 3.9, 0.4, h, 13, PRIMARY if hl else MUTED, bold=True, font=MONO)
    box(s, x, 2.6, 3.9, 1.3, who, 15, TEXT, bold=hl, spacing=1.25)
    hairline(s, x, 3.95, 3.7)
    box(s, x, 4.15, 3.9, 1.2, what, 14, MUTED, spacing=1.25)
    x += 4.15
source(s, "Landscape from idealab #586938 competitor analysis · full matrix on the landing page")

# ── 10 · CLOSING (dark) ─────────────────────────────────────────
s = slide(dark=True)
box(s, 0.7, 1.3, 8, 0.4, "RUN IT YOURSELF", 13, ACCENT, bold=True, font=MONO)
box(s, 0.7, 1.9, 11.9, 1.1, "git clone · make data · make train · make api", 34,
    RGBColor(0xF4, 0xF5, 0xF1), bold=True)
hairline(s, 0.7, 3.3, 11.9, RGBColor(0x2A, 0x2F, 0x2B))
box(s, 0.7, 3.6, 11.5, 1.6,
    "Or just:  ./start.sh   — one script boots the API, dashboards, live feeds\n"
    "and the AI analyst on your machine. GPU training included.",
    20, RGBColor(0xE7, 0xEC, 0xE6), spacing=1.4)
box(s, 0.7, 5.6, 11.9, 0.8,
    "github.com/devansh0703/voyance2  ·  Apache-2.0  ·  built on open data only",
    14, RGBColor(0x8B, 0x94, 0x8C), font=MONO)
box(s, 0.7, 6.5, 11.9, 0.6, "Voyance II — open-source maritime voyage intelligence", 16,
    ACCENT, bold=True)

prs.save("/home/devansh/voyance2/reports/voyance2-deck.pptx")
print("deck saved:", len(prs.slides.__iter__.__self__._sldIdLst), "slides")
